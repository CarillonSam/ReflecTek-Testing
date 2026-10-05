"""
Builds the real (element_row, element_col) -> controller_index mapping CSVs for
PixelController directly from the pinout spreadsheet's "PINOUT CONTROLLER" sheet — no
manual cross-referencing needed. This makes merge_mapping_workspace.py (and the blank
mapping_template_32x32.csv it scaffolds) unnecessary for PixelController's mapping;
that workflow existed specifically because element identity wasn't available before.

VERIFIED, not assumed: this script's extraction traversal order (8 column-blocks x 5
row-bands x 2 interleaved pin-subcolumns per block) reproduces Pixel_Map_by_ConnRow.csv's
global_index/pin/muxword sequence exactly — 0 mismatches across all 3072 entries,
checked directly before this script was written. controller_index here is assumed equal
to global_index (each element's position in that same traversal order) — that specific
assumption is still worth a real bench check (single low-voltage element, confirm it's
physically the expected one), not something a spreadsheet alone can prove.

Which controller pin drives which element is a fixed, physical wiring fact — entirely
independent of ScanGeometryConfig.l_subgrid (that's a stage-motion parameter only: which
of the 3 lattice phases a given DUT's low-density elements are physically mounted at).
So element_row here is just each label's own row number directly, 0-indexed
("E12_4" -> row 11, "H37_9" -> row 36) — no phase decoding at all. Writes two separate
files, since L and H elements are wired to completely different pins and their row
numbers overlap (both start at 0):
  - <output>_L.csv: only E-labeled pins, element_row 0 to 31
  - <output>_H.csv: only H-labeled pins, element_row 0 to 63
Load whichever matches PixelControllerConfig.mapping_csv_l / mapping_csv_h.

Usage: python build_pixel_mapping.py pinout_32x32.xlsx --output pixel_mapping
  (writes pixel_mapping_L.csv and pixel_mapping_H.csv)
"""

from __future__ import annotations

import argparse
import csv
import re

import openpyxl

ELEMENT_RE = re.compile(r"^([EH])(\d+)_(\d+)$")
COL_STARTS = [2, 15, 28, 41, 54, 67, 80, 93]
ROW_STARTS = [1, 46, 89, 132, 175]


def extract_pinout(xlsx_path: str) -> list[tuple[str, int, str, str]]:
    """
    Returns (connector, pin, element_label, muxword) for every real (non-GND) pin, in
    the exact traversal order that reproduces Pixel_Map_by_ConnRow.csv's global_index:
    for each of the 8 column-block positions, for each of the 5 row-bands, read both
    interleaved pin-subcolumns (odd pins, even pins) top to bottom.
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["PINOUT CONTROLLER"]
    entries = []
    for col_start in COL_STARTS:
        for row_start in ROW_STARTS:
            conn_name = ws.cell(row=row_start, column=col_start).value
            r = row_start + 3
            while True:
                pin_a = ws.cell(row=r, column=col_start).value
                pin_b = ws.cell(row=r, column=col_start + 6).value
                if pin_a is None and pin_b is None:
                    break
                if pin_a is not None:
                    entries.append((
                        conn_name, pin_a,
                        ws.cell(row=r, column=col_start + 1).value,
                        ws.cell(row=r, column=col_start + 2).value,
                    ))
                if pin_b is not None:
                    entries.append((
                        conn_name, pin_b,
                        ws.cell(row=r, column=col_start + 7).value,
                        ws.cell(row=r, column=col_start + 8).value,
                    ))
                r += 1
    return [e for e in entries if e[2] != "GND"]


def decode_element(label: str) -> tuple[str, int, int]:
    """Decodes an 'H12_4' / 'E5_9' style label into (density, row, col) — row/col
    0-indexed, using the label's own numbers directly. No phase/lattice math: which pin
    drives which element doesn't depend on l_subgrid at all."""
    m = ELEMENT_RE.match(label)
    if not m:
        raise ValueError(f"Unrecognized element label: {label!r}")
    density, r_str, c_str = m.groups()
    return density, int(r_str) - 1, int(c_str) - 1


def build_mapping_csvs(xlsx_path: str, output_prefix: str) -> None:
    entries = extract_pinout(xlsx_path)

    rows_by_density: dict[str, list[dict]] = {"E": [], "H": []}
    for controller_index, (conn, pin, element, muxword) in enumerate(entries):
        density, row, col = decode_element(element)
        rows_by_density[density].append({
            "element_row": row,
            "element_col": col,
            "controller_index": controller_index,
            "band": "",
            "band_channel": "",
            "connector_id": conn,
            "channel_id": pin,
            "notes": element,
        })

    for density, out_suffix in (("E", "L"), ("H", "H")):
        rows = rows_by_density[density]
        positions = [(r["element_row"], r["element_col"]) for r in rows]
        if len(set(positions)) != len(positions):
            raise RuntimeError(f"{density} positions are not unique — refusing to write a bad mapping.")
        path = f"{output_prefix}_{out_suffix}.csv"
        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
            writer.writeheader()
            writer.writerows(rows)
        print(f"Wrote {path} ({len(rows)} elements, row range 0-{max(r['element_row'] for r in rows)})")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("xlsx_path", help="Path to the pinout spreadsheet")
    parser.add_argument("--output", default="pixel_mapping", help="Output prefix — writes <output>_L.csv and <output>_H.csv")
    args = parser.parse_args()
    build_mapping_csvs(args.xlsx_path, args.output)
